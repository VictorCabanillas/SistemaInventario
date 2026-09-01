from fastapi import FastAPI, HTTPException, Query, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, StreamingResponse
from typing import List, Optional
import os
import io
import re
import shutil
import sqlite3
import tempfile

ICONOS_DIR = os.path.join(os.path.dirname(__file__), "iconos")

from database import get_db, init_db, DB_PATH
from schemas import (
    Categoria, CategoriaCreate,
    Sala, SalaCreate,
    Operador, OperadorCreate,
    Proyecto, ProyectoCreate, ProyectoUpdate,
    Articulo, ArticuloCreate, ArticuloUpdate, StockUpdate,
    ArticuloDetalle, StockPorProyecto, StockLinea,
    Movimiento, MovimientoCreate,
    ResultadoBusqueda, Sugerencias, CatalogoItem, FusionArticulos
)

app = FastAPI(title="Inventario Almacén", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

# Inicializar DB al arrancar
@app.on_event("startup")
def startup():
    init_db()


def _find_or_create_stock(db, articulo_id, proyecto_id, ubicacion, estado='ok'):
    """Busca una fila de stock por (articulo, proyecto, ubicación normalizada,
    estado); si no existe, la crea con cantidad 0. Devuelve su id."""
    row = db.execute(
        "SELECT id FROM stock WHERE articulo_id=? AND proyecto_id=? "
        "AND COALESCE(ubicacion,'')=COALESCE(?,'') AND estado=?",
        (articulo_id, proyecto_id, ubicacion, estado)
    ).fetchone()
    if row:
        return row["id"]
    cur = db.execute(
        "INSERT INTO stock (articulo_id, proyecto_id, cantidad, ubicacion, estado) VALUES (?,?,0,?,?)",
        (articulo_id, proyecto_id, ubicacion, estado)
    )
    return cur.lastrowid


def _crear_o_sumar_articulo(db, *, nombre, marca, referencia, categoria_id, unidad,
                             stock_minimo, sala_id, ubicacion, notas, proyecto_id, cantidad,
                             motivo="Alta inicial", operador="Sistema"):
    """Crea el artículo de catálogo si no existe (por nombre+marca+referencia,
    con marca/referencia normalizadas a '' si vienen vacías) y suma `cantidad`
    a su fila de stock en (proyecto_id, ubicación). Devuelve (articulo_id, stock_id)."""
    marca = marca or ""
    referencia = referencia or ""

    articulo = db.execute(
        "SELECT id FROM articulos WHERE nombre=? AND marca=? AND referencia=?",
        (nombre, marca, referencia)
    ).fetchone()
    if articulo:
        articulo_id = articulo["id"]
    else:
        cur = db.execute("""
            INSERT INTO articulos (nombre, marca, referencia, categoria_id, unidad, stock_minimo, sala_id, ubicacion, notas)
            VALUES (?,?,?,?,?,?,?,?,?)
        """, (nombre, marca, referencia, categoria_id, unidad, stock_minimo, sala_id, ubicacion, notas))
        articulo_id = cur.lastrowid

    stock_id = _find_or_create_stock(db, articulo_id, proyecto_id, ubicacion, estado='ok')
    if sala_id is not None:
        db.execute("UPDATE stock SET sala_id=? WHERE id=?", (sala_id, stock_id))
    if cantidad:
        db.execute(
            "UPDATE stock SET cantidad = cantidad + ?, updated_at=datetime('now') WHERE id=?",
            (cantidad, stock_id)
        )
        db.execute("""
            INSERT INTO movimientos (articulo_id, tipo, cantidad, motivo, operador, proyecto_id)
            VALUES (?, 'entrada', ?, ?, ?, ?)
        """, (articulo_id, cantidad, motivo, operador, proyecto_id))

    return articulo_id, stock_id


def _restar_stock(db, stock_id, cantidad):
    """Resta `cantidad` de una fila de stock; si queda en ~0, borra la fila
    (y el artículo de catálogo si ya no tiene stock en ningún sitio) en vez
    de dejar una fila fantasma con cantidad 0."""
    db.execute(
        "UPDATE stock SET cantidad = cantidad - ?, updated_at=datetime('now') WHERE id=?",
        (cantidad, stock_id)
    )
    row = db.execute("SELECT articulo_id, cantidad FROM stock WHERE id=?", (stock_id,)).fetchone()
    if row and abs(row["cantidad"]) < 1e-9:
        db.execute("DELETE FROM stock WHERE id=?", (stock_id,))
        otros = db.execute("SELECT COUNT(*) as n FROM stock WHERE articulo_id=?", (row["articulo_id"],)).fetchone()
        if otros["n"] == 0:
            db.execute("DELETE FROM articulos WHERE id=?", (row["articulo_id"],))


# ─────────────────────────────────────────
# ICONOS PERSONALIZADOS
# ─────────────────────────────────────────

ICONOS_MIME = {
    ".svg": "image/svg+xml",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
}

@app.get("/api/iconos")
def get_iconos():
    os.makedirs(ICONOS_DIR, exist_ok=True)
    return sorted(
        f for f in os.listdir(ICONOS_DIR)
        if os.path.splitext(f)[1].lower() in ICONOS_MIME
    )

@app.get("/api/iconos/{filename}")
def get_icono(filename: str):
    if "/" in filename or "\\" in filename:
        raise HTTPException(400, "Nombre de archivo inválido")
    ext = os.path.splitext(filename)[1].lower()
    if ext not in ICONOS_MIME:
        raise HTTPException(400, "Tipo de archivo no permitido")
    path = os.path.join(ICONOS_DIR, filename)
    if not os.path.exists(path):
        raise HTTPException(404, "Icono no encontrado")
    return FileResponse(path, media_type=ICONOS_MIME[ext])

@app.delete("/api/iconos/{filename}", status_code=204)
def delete_icono(filename: str):
    if "/" in filename or "\\" in filename:
        raise HTTPException(400, "Nombre de archivo inválido")
    path = os.path.join(ICONOS_DIR, filename)
    if os.path.exists(path):
        os.remove(path)


ICONO_MAX_UPLOAD_BYTES = 5 * 1024 * 1024  # 5 MB
ICONO_LIENZO_PX = 256  # tamaño del icono procesado (cuadrado, con transparencia)


def _slugify(nombre: str) -> str:
    base = re.sub(r"[^a-zA-Z0-9]+", "-", nombre).strip("-").lower()
    return base or "icono"


def _nombre_disponible(base: str, ext: str) -> str:
    nombre = f"{base}{ext}"
    i = 2
    while os.path.exists(os.path.join(ICONOS_DIR, nombre)):
        nombre = f"{base}-{i}{ext}"
        i += 1
    return nombre


@app.post("/api/iconos", status_code=201)
async def subir_icono(file: UploadFile = File(...)):
    """Sube una imagen y la procesa para usarla como icono de proyecto:
    SVG se guarda tal cual (ya es apto como máscara); el resto se normaliza
    a PNG, se ajusta (sin deformar) a un lienzo cuadrado transparente y se
    redimensiona a un tamaño manejable."""
    os.makedirs(ICONOS_DIR, exist_ok=True)
    content = await file.read()
    if not content:
        raise HTTPException(400, "El archivo está vacío")
    if len(content) > ICONO_MAX_UPLOAD_BYTES:
        raise HTTPException(400, "La imagen no puede superar los 5 MB")

    nombre_original = file.filename or "icono"
    ext_original = os.path.splitext(nombre_original)[1].lower()
    base = _slugify(os.path.splitext(nombre_original)[0])

    if ext_original == ".svg":
        texto = content.decode("utf-8", errors="ignore")
        if "<svg" not in texto.lower():
            raise HTTPException(400, "El archivo SVG no es válido")
        filename = _nombre_disponible(base, ".svg")
        with open(os.path.join(ICONOS_DIR, filename), "wb") as f:
            f.write(content)
        return {"filename": filename}

    try:
        from PIL import Image, UnidentifiedImageError
    except ImportError:
        raise HTTPException(500, "Pillow no instalado en el servidor")

    try:
        img = Image.open(io.BytesIO(content))
        img.load()
    except (UnidentifiedImageError, OSError):
        raise HTTPException(400, "El archivo no es una imagen válida")

    img = img.convert("RGBA")
    img.thumbnail((ICONO_LIENZO_PX, ICONO_LIENZO_PX), Image.LANCZOS)
    lienzo = Image.new("RGBA", (ICONO_LIENZO_PX, ICONO_LIENZO_PX), (0, 0, 0, 0))
    offset = ((ICONO_LIENZO_PX - img.width) // 2, (ICONO_LIENZO_PX - img.height) // 2)
    lienzo.paste(img, offset, img)

    filename = _nombre_disponible(base, ".png")
    lienzo.save(os.path.join(ICONOS_DIR, filename), "PNG")
    return {"filename": filename}


# ─────────────────────────────────────────
# CATEGORÍAS
# ─────────────────────────────────────────

@app.get("/api/categorias", response_model=List[Categoria])
def get_categorias():
    db = get_db()
    rows = db.execute("SELECT id, nombre FROM categorias ORDER BY nombre").fetchall()
    db.close()
    return [dict(r) for r in rows]

@app.post("/api/categorias", response_model=Categoria, status_code=201)
def create_categoria(data: CategoriaCreate):
    db = get_db()
    try:
        cur = db.execute("INSERT INTO categorias (nombre) VALUES (?)", (data.nombre,))
        db.commit()
        row = db.execute("SELECT id, nombre FROM categorias WHERE id=?", (cur.lastrowid,)).fetchone()
        return dict(row)
    except Exception as e:
        raise HTTPException(400, f"Categoría ya existe: {e}")
    finally:
        db.close()

@app.delete("/api/categorias/{id}", status_code=204)
def delete_categoria(id: int):
    db = get_db()
    db.execute("DELETE FROM categorias WHERE id=?", (id,))
    db.commit()
    db.close()


# ─────────────────────────────────────────
# SALAS
# ─────────────────────────────────────────

@app.get("/api/salas", response_model=List[Sala])
def get_salas():
    db = get_db()
    rows = db.execute("SELECT id, nombre FROM salas ORDER BY nombre").fetchall()
    db.close()
    return [dict(r) for r in rows]

@app.post("/api/salas", response_model=Sala, status_code=201)
def create_sala(data: SalaCreate):
    db = get_db()
    try:
        cur = db.execute("INSERT INTO salas (nombre) VALUES (?)", (data.nombre,))
        db.commit()
        row = db.execute("SELECT id, nombre FROM salas WHERE id=?", (cur.lastrowid,)).fetchone()
        return dict(row)
    except Exception as e:
        raise HTTPException(400, f"La sala ya existe: {e}")
    finally:
        db.close()

@app.delete("/api/salas/{id}", status_code=204)
def delete_sala(id: int):
    db = get_db()
    db.execute("DELETE FROM salas WHERE id=?", (id,))
    db.commit()
    db.close()


# ─────────────────────────────────────────
# OPERADORES
# ─────────────────────────────────────────

@app.get("/api/operadores", response_model=List[Operador])
def get_operadores():
    db = get_db()
    rows = db.execute("SELECT id, nombre FROM operadores ORDER BY nombre").fetchall()
    db.close()
    return [dict(r) for r in rows]

@app.post("/api/operadores", response_model=Operador, status_code=201)
def create_operador(data: OperadorCreate):
    db = get_db()
    try:
        cur = db.execute("INSERT INTO operadores (nombre) VALUES (?)", (data.nombre,))
        db.commit()
        row = db.execute("SELECT id, nombre FROM operadores WHERE id=?", (cur.lastrowid,)).fetchone()
        return dict(row)
    except Exception as e:
        raise HTTPException(400, f"El operador ya existe: {e}")
    finally:
        db.close()

@app.delete("/api/operadores/{id}", status_code=204)
def delete_operador(id: int):
    db = get_db()
    db.execute("DELETE FROM operadores WHERE id=?", (id,))
    db.commit()
    db.close()


# ─────────────────────────────────────────
# PROYECTOS
# ─────────────────────────────────────────
# total_articulos / articulos_bajo_minimo se calculan con subconsultas
# correlacionadas: el stock mínimo es un valor único por artículo (catálogo),
# así que "bajo mínimo" compara la SUMA de todas las ubicaciones 'ok' de ese
# artículo en el proyecto contra ese mínimo — no la ubicación individual.

PROYECTO_SELECT = """
    SELECT p.*,
        (SELECT COUNT(DISTINCT s.articulo_id) FROM stock s
         WHERE s.proyecto_id = p.id AND s.estado = 'ok') AS total_articulos,
        (SELECT COUNT(*) FROM (
            SELECT s.articulo_id, SUM(s.cantidad) AS total, MAX(a.stock_minimo) AS minimo
            FROM stock s JOIN articulos a ON a.id = s.articulo_id
            WHERE s.proyecto_id = p.id AND s.estado = 'ok' AND a.stock_minimo IS NOT NULL
            GROUP BY s.articulo_id
            HAVING total < minimo
        )) AS articulos_bajo_minimo
    FROM proyectos p
"""

@app.get("/api/proyectos", response_model=List[Proyecto])
def get_proyectos():
    db = get_db()
    rows = db.execute(PROYECTO_SELECT + "ORDER BY p.nombre").fetchall()
    db.close()
    return [dict(r) for r in rows]

@app.post("/api/proyectos", response_model=Proyecto, status_code=201)
def create_proyecto(data: ProyectoCreate):
    db = get_db()
    cur = db.execute(
        "INSERT INTO proyectos (nombre, descripcion, color, icono) VALUES (?,?,?,?)",
        (data.nombre, data.descripcion, data.color, data.icono)
    )
    db.commit()
    row = db.execute("""
        SELECT p.*, 0 as total_articulos, 0 as articulos_bajo_minimo
        FROM proyectos p WHERE p.id=?
    """, (cur.lastrowid,)).fetchone()
    db.close()
    return dict(row)

@app.put("/api/proyectos/{id}", response_model=Proyecto)
def update_proyecto(id: int, data: ProyectoUpdate):
    db = get_db()
    fields = {k: v for k, v in data.dict().items() if v is not None}
    if fields:
        sets = ", ".join(f"{k}=?" for k in fields)
        db.execute(f"UPDATE proyectos SET {sets} WHERE id=?", (*fields.values(), id))
        db.commit()
    row = db.execute(PROYECTO_SELECT + "WHERE p.id=?", (id,)).fetchone()
    db.close()
    if not row:
        raise HTTPException(404, "Proyecto no encontrado")
    return dict(row)

@app.delete("/api/proyectos/{id}", status_code=204)
def delete_proyecto(id: int):
    db = get_db()
    row = db.execute("SELECT es_almacen FROM proyectos WHERE id=?", (id,)).fetchone()
    if row and row["es_almacen"]:
        db.close()
        raise HTTPException(400, "No se puede eliminar el proyecto Almacén general")
    db.execute("DELETE FROM proyectos WHERE id=?", (id,))
    db.commit()
    db.close()


# ─────────────────────────────────────────
# ARTÍCULOS
#
# `articulos` es el catálogo global (único por nombre+marca+referencia; el
# stock_mínimo también vive aquí, es un valor único por artículo).
# `stock` guarda las existencias de un artículo en un proyecto concreto,
# repartidas en filas por ubicación/estado ('ok' o 'baja'). Los listados
# (ARTICULO_LISTADO_SELECT) agregan todas las ubicaciones 'ok' de cada
# artículo+proyecto en una sola fila con el total; el detalle de un artículo
# devuelve ese total más el desglose fila a fila.
# ─────────────────────────────────────────

ARTICULO_LISTADO_SELECT = """
    SELECT
        a.id AS id,
        s.proyecto_id AS proyecto_id,
        p.nombre AS proyecto_nombre,
        p.color AS proyecto_color,
        a.nombre AS nombre,
        a.marca AS marca,
        a.referencia AS referencia,
        a.categoria_id AS categoria_id,
        c.nombre AS categoria_nombre,
        a.unidad AS unidad,
        SUM(s.cantidad) AS cantidad,
        a.stock_minimo AS stock_minimo,
        CASE WHEN a.stock_minimo IS NOT NULL AND SUM(s.cantidad) < a.stock_minimo
             THEN 1 ELSE 0 END AS bajo_minimo,
        COUNT(*) AS num_ubicaciones,
        CASE WHEN COUNT(*) = 1 THEN MAX(s.ubicacion) END AS ubicacion,
        CASE WHEN COUNT(*) = 1 THEN MAX(sal.nombre) END AS sala_nombre,
        CASE WHEN SUM(CASE WHEN s.ubicacion IS NOT NULL OR s.sala_id IS NOT NULL THEN 1 ELSE 0 END) = 0
             THEN 1 ELSE 0 END AS sin_ubicacion,
        a.notas AS notas,
        MIN(s.created_at) AS created_at,
        MAX(s.updated_at) AS updated_at
    FROM stock s
    JOIN articulos a ON a.id = s.articulo_id
    JOIN proyectos p ON p.id = s.proyecto_id
    LEFT JOIN categorias c ON c.id = a.categoria_id
    LEFT JOIN salas sal ON sal.id = s.sala_id
    WHERE s.estado = 'ok'
"""

# Fila plana (sin agregar), una por ubicación: usada por el listado global de
# bajas y por el ajuste masivo, que necesitan granularidad por ubicación.
STOCK_LINEA_SELECT = """
    SELECT
        s.id AS stock_id,
        a.id AS articulo_id,
        a.nombre AS nombre,
        a.marca AS marca,
        a.referencia AS referencia,
        a.unidad AS unidad,
        s.cantidad AS cantidad,
        s.ubicacion AS ubicacion,
        sal.nombre AS sala_nombre,
        s.estado AS estado,
        s.proyecto_id AS proyecto_id,
        p.nombre AS proyecto_nombre,
        p.color AS proyecto_color,
        s.updated_at AS updated_at
    FROM stock s
    JOIN articulos a ON a.id = s.articulo_id
    JOIN proyectos p ON p.id = s.proyecto_id
    LEFT JOIN salas sal ON sal.id = s.sala_id
"""


# Filtra por artículos que tengan ALGUNA ubicación en la sala indicada, sin
# restringir la agregación: el total mostrado sigue siendo la suma de TODAS
# sus ubicaciones (no solo la de esa sala).
SALA_EXISTS_CLAUSE = """AND EXISTS (
    SELECT 1 FROM stock s2
    WHERE s2.articulo_id = a.id AND s2.proyecto_id = s.proyecto_id
      AND s2.sala_id = ? AND s2.estado = 'ok'
) """

@app.get("/api/articulos", response_model=List[Articulo])
def get_todos_articulos(sala_id: Optional[int] = None):
    """Todos los artículos con stock 'ok', de todos los proyectos (vista del Almacén general)."""
    db = get_db()
    query = ARTICULO_LISTADO_SELECT
    params = []
    if sala_id is not None:
        query += SALA_EXISTS_CLAUSE
        params.append(sala_id)
    query += "GROUP BY a.id, s.proyecto_id ORDER BY a.nombre"
    rows = db.execute(query, params).fetchall()
    db.close()
    return [dict(r) for r in rows]

@app.get("/api/proyectos/{proyecto_id}/articulos", response_model=List[Articulo])
def get_articulos(proyecto_id: int, sala_id: Optional[int] = None):
    db = get_db()
    query = ARTICULO_LISTADO_SELECT + "AND s.proyecto_id = ? "
    params = [proyecto_id]
    if sala_id is not None:
        query += SALA_EXISTS_CLAUSE
        params.append(sala_id)
    query += "GROUP BY a.id, s.proyecto_id ORDER BY a.nombre"
    rows = db.execute(query, params).fetchall()
    db.close()
    return [dict(r) for r in rows]

@app.get("/api/alertas", response_model=List[Articulo])
def get_alertas():
    """Artículos cuyo total agregado (todas las ubicaciones 'ok') está por debajo de su mínimo."""
    db = get_db()
    rows = db.execute(
        ARTICULO_LISTADO_SELECT + "GROUP BY a.id, s.proyecto_id HAVING bajo_minimo = 1 ORDER BY a.nombre"
    ).fetchall()
    db.close()
    return [dict(r) for r in rows]

@app.get("/api/bajas", response_model=List[StockLinea])
def get_bajas():
    """Todas las filas de stock marcadas como 'baja' (material no funcional), de cualquier proyecto."""
    db = get_db()
    rows = db.execute(STOCK_LINEA_SELECT + "WHERE s.estado = 'baja' ORDER BY a.nombre").fetchall()
    db.close()
    return [dict(r) for r in rows]

@app.get("/api/proyectos/{proyecto_id}/stock", response_model=List[StockLinea])
def get_stock_proyecto(proyecto_id: int):
    """Filas de stock 'ok' en bruto (una por ubicación) de un proyecto — usado por el ajuste masivo."""
    db = get_db()
    rows = db.execute(
        STOCK_LINEA_SELECT + "WHERE s.proyecto_id = ? AND s.estado = 'ok' ORDER BY a.nombre, s.ubicacion",
        (proyecto_id,)
    ).fetchall()
    db.close()
    return [dict(r) for r in rows]

@app.get("/api/stock", response_model=List[StockLinea])
def get_stock_global():
    """Todas las filas de stock 'ok' en bruto, de todos los proyectos — usado por
    el ajuste masivo en la vista de Almacén general (que agrega todos los proyectos)."""
    db = get_db()
    rows = db.execute(STOCK_LINEA_SELECT + "WHERE s.estado = 'ok' ORDER BY a.nombre, s.ubicacion").fetchall()
    db.close()
    return [dict(r) for r in rows]

@app.get("/api/proyectos/{proyecto_id}/articulos/{articulo_id}", response_model=ArticuloDetalle)
def get_articulo_detalle(proyecto_id: int, articulo_id: int):
    db = get_db()
    cab = db.execute("""
        SELECT a.id, a.nombre, a.marca, a.referencia, a.categoria_id, c.nombre AS categoria_nombre,
               a.unidad, a.stock_minimo, a.notas,
               p.id AS proyecto_id, p.nombre AS proyecto_nombre, p.color AS proyecto_color
        FROM articulos a
        JOIN proyectos p ON p.id = ?
        LEFT JOIN categorias c ON c.id = a.categoria_id
        WHERE a.id = ?
    """, (proyecto_id, articulo_id)).fetchone()
    if not cab:
        db.close()
        raise HTTPException(404, "Artículo no encontrado")

    filas = db.execute("""
        SELECT s.id AS stock_id, s.ubicacion, s.sala_id, sal.nombre AS sala_nombre,
               s.cantidad, s.estado, s.created_at, s.updated_at
        FROM stock s
        LEFT JOIN salas sal ON sal.id = s.sala_id
        WHERE s.articulo_id = ? AND s.proyecto_id = ?
        ORDER BY s.ubicacion
    """, (articulo_id, proyecto_id)).fetchall()
    db.close()

    ubicaciones = [dict(f) for f in filas if f["estado"] == "ok"]
    bajas = [dict(f) for f in filas if f["estado"] == "baja"]
    cantidad_total = sum(u["cantidad"] for u in ubicaciones)
    minimo = cab["stock_minimo"]
    bajo_minimo = minimo is not None and cantidad_total < minimo

    return {
        **dict(cab),
        "cantidad_total": cantidad_total,
        "bajo_minimo": bajo_minimo,
        "ubicaciones": ubicaciones,
        "bajas": bajas,
    }

@app.get("/api/articulos/{articulo_id}/stock", response_model=List[StockPorProyecto])
def get_stock_por_articulo(articulo_id: int):
    """Total agregado de este artículo en cada proyecto donde tiene stock 'ok' — 'también en otros proyectos'."""
    db = get_db()
    rows = db.execute("""
        SELECT s.proyecto_id AS proyecto_id, p.nombre AS proyecto_nombre, p.color AS proyecto_color,
               SUM(s.cantidad) AS cantidad_total
        FROM stock s
        JOIN proyectos p ON p.id = s.proyecto_id
        WHERE s.articulo_id = ? AND s.estado = 'ok'
        GROUP BY s.proyecto_id
        ORDER BY p.nombre
    """, (articulo_id,)).fetchall()
    db.close()
    return [dict(r) for r in rows]

@app.post("/api/articulos", status_code=201)
def create_articulo(data: ArticuloCreate):
    db = get_db()
    try:
        if data.proyecto_id is not None:
            proyecto = db.execute("SELECT id FROM proyectos WHERE id=?", (data.proyecto_id,)).fetchone()
            if not proyecto:
                raise HTTPException(404, "Proyecto no encontrado")
            proyecto_id = data.proyecto_id
        else:
            almacen = db.execute("SELECT id FROM proyectos WHERE es_almacen = 1").fetchone()
            if not almacen:
                raise HTTPException(500, "No existe un proyecto Almacén general")
            proyecto_id = almacen["id"]

        articulo_id, stock_id = _crear_o_sumar_articulo(
            db, nombre=data.nombre, marca=data.marca, referencia=data.referencia,
            categoria_id=data.categoria_id, unidad=data.unidad, stock_minimo=data.stock_minimo,
            sala_id=data.sala_id, ubicacion=data.ubicacion, notas=data.notas,
            proyecto_id=proyecto_id, cantidad=data.cantidad,
        )

        db.commit()
        return {"articulo_id": articulo_id, "stock_id": stock_id}
    except sqlite3.IntegrityError as e:
        raise HTTPException(400, f"No se pudo crear el artículo: {e}")
    finally:
        db.close()

@app.put("/api/articulos/{id}")
def update_articulo(id: int, data: ArticuloUpdate):
    """Actualiza solo campos de catálogo (compartidos por todos los proyectos/ubicaciones)."""
    db = get_db()
    fields = {k: v for k, v in data.dict().items() if v is not None}
    if fields:
        sets = ", ".join(f"{k}=?" for k in fields) + ", updated_at=datetime('now')"
        try:
            db.execute(f"UPDATE articulos SET {sets} WHERE id=?", (*fields.values(), id))
            db.commit()
        except sqlite3.IntegrityError as e:
            db.close()
            raise HTTPException(400, f"No se pudo actualizar: {e}")
    row = db.execute("SELECT * FROM articulos WHERE id=?", (id,)).fetchone()
    db.close()
    if not row:
        raise HTTPException(404, "Artículo no encontrado")
    return dict(row)

@app.put("/api/stock/{stock_id}")
def update_stock(stock_id: int, data: StockUpdate):
    """Actualiza ubicación/sala de una fila de stock concreta. Si la nueva
    ubicación coincide con otra fila existente del mismo artículo+proyecto+
    estado, las fusiona sumando cantidades."""
    db = get_db()
    fields = {k: v for k, v in data.dict().items() if v is not None}
    if not fields:
        db.close()
        return {"detail": "sin cambios"}

    origen = db.execute("SELECT * FROM stock WHERE id=?", (stock_id,)).fetchone()
    if not origen:
        db.close()
        raise HTTPException(404, "Stock no encontrado")

    sets = ", ".join(f"{k}=?" for k in fields) + ", updated_at=datetime('now')"
    try:
        db.execute(f"UPDATE stock SET {sets} WHERE id=?", (*fields.values(), stock_id))
        db.commit()
        db.close()
        return {"detail": "ok", "fusionado": False, "stock_id": stock_id}
    except sqlite3.IntegrityError:
        nueva_ubicacion = fields.get("ubicacion", origen["ubicacion"])
        destino = db.execute(
            "SELECT * FROM stock WHERE articulo_id=? AND proyecto_id=? "
            "AND COALESCE(ubicacion,'')=COALESCE(?,'') AND estado=? AND id != ?",
            (origen["articulo_id"], origen["proyecto_id"], nueva_ubicacion, origen["estado"], stock_id)
        ).fetchone()
        if not destino:
            db.close()
            raise HTTPException(400, "No se pudo actualizar la ubicación")
        db.execute(
            "UPDATE stock SET cantidad = cantidad + ?, updated_at=datetime('now') WHERE id=?",
            (origen["cantidad"], destino["id"])
        )
        db.execute("DELETE FROM stock WHERE id=?", (stock_id,))
        db.commit()
        db.close()
        return {"detail": "ok", "fusionado": True, "stock_id": destino["id"]}

@app.delete("/api/proyectos/{proyecto_id}/articulos/{articulo_id}", status_code=204)
def delete_articulo(proyecto_id: int, articulo_id: int):
    """Elimina TODAS las ubicaciones de este artículo en este proyecto."""
    db = get_db()
    db.execute("DELETE FROM stock WHERE articulo_id=? AND proyecto_id=?", (articulo_id, proyecto_id))
    otros = db.execute("SELECT COUNT(*) as n FROM stock WHERE articulo_id=?", (articulo_id,)).fetchone()
    if otros["n"] == 0:
        db.execute("DELETE FROM articulos WHERE id=?", (articulo_id,))
    db.commit()
    db.close()

@app.delete("/api/stock/{stock_id}", status_code=204)
def delete_stock(stock_id: int):
    """Elimina una sola ubicación (fila de stock)."""
    db = get_db()
    row = db.execute("SELECT articulo_id FROM stock WHERE id=?", (stock_id,)).fetchone()
    if row:
        db.execute("DELETE FROM stock WHERE id=?", (stock_id,))
        otros = db.execute("SELECT COUNT(*) as n FROM stock WHERE articulo_id=?", (row["articulo_id"],)).fetchone()
        if otros["n"] == 0:
            db.execute("DELETE FROM articulos WHERE id=?", (row["articulo_id"],))
        db.commit()
    db.close()


# ─────────────────────────────────────────
# MOVIMIENTOS DE STOCK
# ─────────────────────────────────────────

MOVIMIENTO_SELECT = """
    SELECT m.*, a.nombre as articulo_nombre,
           po.nombre as proyecto_nombre,
           porig.nombre as proyecto_origen_nombre,
           pdest.nombre as proyecto_destino_nombre
    FROM movimientos m
    JOIN articulos a ON a.id = m.articulo_id
    LEFT JOIN proyectos po ON po.id = m.proyecto_id
    LEFT JOIN proyectos porig ON porig.id = m.proyecto_origen_id
    LEFT JOIN proyectos pdest ON pdest.id = m.proyecto_destino_id
"""

@app.get("/api/articulos/{articulo_id}/movimientos", response_model=List[Movimiento])
def get_movimientos(articulo_id: int, limit: int = 50):
    db = get_db()
    rows = db.execute(
        MOVIMIENTO_SELECT + "WHERE m.articulo_id = ? ORDER BY m.fecha DESC LIMIT ?",
        (articulo_id, limit)
    ).fetchall()
    db.close()
    return [dict(r) for r in rows]

@app.get("/api/movimientos", response_model=List[Movimiento])
def get_all_movimientos(limit: int = 100):
    db = get_db()
    rows = db.execute(MOVIMIENTO_SELECT + "ORDER BY m.fecha DESC LIMIT ?", (limit,)).fetchall()
    db.close()
    return [dict(r) for r in rows]

@app.post("/api/movimientos", status_code=201)
def create_movimiento(data: MovimientoCreate):
    db = get_db()
    try:
        if data.tipo == "entrada":
            if data.stock_id:
                stock = db.execute("SELECT * FROM stock WHERE id=?", (data.stock_id,)).fetchone()
                if not stock:
                    raise HTTPException(404, "Stock no encontrado")
                articulo_id, proyecto_id, stock_id = stock["articulo_id"], stock["proyecto_id"], stock["id"]
            elif data.articulo_id and data.proyecto_id:
                articulo_id, proyecto_id = data.articulo_id, data.proyecto_id
                stock_id = _find_or_create_stock(db, articulo_id, proyecto_id, data.ubicacion, estado='ok')
                if data.sala_id is not None:
                    db.execute("UPDATE stock SET sala_id=? WHERE id=?", (data.sala_id, stock_id))
            else:
                raise HTTPException(400, "La entrada requiere stock_id, o articulo_id + proyecto_id")

            db.execute(
                "UPDATE stock SET cantidad = cantidad + ?, updated_at=datetime('now') WHERE id=?",
                (data.cantidad, stock_id)
            )
            db.execute("""
                INSERT INTO movimientos (articulo_id, tipo, cantidad, motivo, operador, proyecto_id)
                VALUES (?, 'entrada', ?, ?, ?, ?)
            """, (articulo_id, data.cantidad, data.motivo, data.operador, proyecto_id))
            db.commit()
            return {"detail": "ok", "stock_id": stock_id}

        if data.tipo == "salida":
            if not data.stock_id:
                raise HTTPException(400, "La salida requiere stock_id")
            stock = db.execute("SELECT * FROM stock WHERE id=?", (data.stock_id,)).fetchone()
            if not stock:
                raise HTTPException(404, "Stock no encontrado")
            if stock["cantidad"] < data.cantidad:
                raise HTTPException(400, f"Stock insuficiente. Stock actual: {stock['cantidad']}")

            _restar_stock(db, data.stock_id, data.cantidad)
            db.execute("""
                INSERT INTO movimientos (articulo_id, tipo, cantidad, motivo, operador, proyecto_id)
                VALUES (?, 'salida', ?, ?, ?, ?)
            """, (stock["articulo_id"], data.cantidad, data.motivo, data.operador, stock["proyecto_id"]))
            db.commit()
            return {"detail": "ok"}

        if data.tipo in ("baja", "reparacion"):
            if not data.stock_id:
                raise HTTPException(400, f"El movimiento de {data.tipo} requiere stock_id")
            origen = db.execute("SELECT * FROM stock WHERE id=?", (data.stock_id,)).fetchone()
            if not origen:
                raise HTTPException(404, "Stock no encontrado")

            estado_origen_esperado = "ok" if data.tipo == "baja" else "baja"
            estado_destino = "baja" if data.tipo == "baja" else "ok"
            if origen["estado"] != estado_origen_esperado:
                raise HTTPException(400, f"El stock no está en estado '{estado_origen_esperado}'")
            if origen["cantidad"] < data.cantidad:
                raise HTTPException(400, f"Cantidad insuficiente. Disponible: {origen['cantidad']}")

            ubicacion_destino = data.ubicacion_destino if data.ubicacion_destino is not None else origen["ubicacion"]
            destino_id = _find_or_create_stock(
                db, origen["articulo_id"], origen["proyecto_id"], ubicacion_destino, estado=estado_destino
            )
            if data.sala_destino_id is not None:
                db.execute("UPDATE stock SET sala_id=? WHERE id=?", (data.sala_destino_id, destino_id))

            _restar_stock(db, origen["id"], data.cantidad)
            db.execute(
                "UPDATE stock SET cantidad = cantidad + ?, updated_at=datetime('now') WHERE id=?",
                (data.cantidad, destino_id)
            )
            db.execute("""
                INSERT INTO movimientos (articulo_id, tipo, cantidad, motivo, operador, proyecto_id)
                VALUES (?,?,?,?,?,?)
            """, (origen["articulo_id"], data.tipo, data.cantidad, data.motivo, data.operador, origen["proyecto_id"]))
            db.commit()
            return {"detail": "ok"}

        if data.tipo == "transferencia":
            if not (data.stock_id and data.proyecto_destino_id):
                raise HTTPException(400, "La transferencia requiere stock_id y proyecto_destino_id")

            origen = db.execute("SELECT * FROM stock WHERE id=?", (data.stock_id,)).fetchone()
            if not origen:
                raise HTTPException(404, "Stock de origen no encontrado")
            if origen["proyecto_id"] == data.proyecto_destino_id:
                raise HTTPException(400, "El proyecto de origen y el de destino deben ser distintos")
            if origen["estado"] != "ok":
                raise HTTPException(400, "No se puede transferir material dado de baja; repáralo primero")
            if origen["cantidad"] < data.cantidad:
                raise HTTPException(400, f"Stock insuficiente en el proyecto de origen. Disponible: {origen['cantidad']}")

            ubicacion_destino = data.ubicacion_destino if data.ubicacion_destino is not None else origen["ubicacion"]
            destino_id = _find_or_create_stock(
                db, origen["articulo_id"], data.proyecto_destino_id, ubicacion_destino, estado='ok'
            )

            _restar_stock(db, origen["id"], data.cantidad)
            db.execute(
                "UPDATE stock SET cantidad = cantidad + ?, updated_at=datetime('now') WHERE id=?",
                (data.cantidad, destino_id)
            )
            db.execute("""
                INSERT INTO movimientos (articulo_id, tipo, cantidad, motivo, operador, proyecto_origen_id, proyecto_destino_id)
                VALUES (?, 'transferencia', ?, ?, ?, ?, ?)
            """, (origen["articulo_id"], data.cantidad, data.motivo, data.operador,
                  origen["proyecto_id"], data.proyecto_destino_id))
            db.commit()
            return {"detail": "ok"}

        raise HTTPException(400, f"Tipo de movimiento desconocido: {data.tipo}")
    finally:
        db.close()


# ─────────────────────────────────────────
# BÚSQUEDA GLOBAL
# ─────────────────────────────────────────

@app.get("/api/buscar", response_model=List[ResultadoBusqueda])
def buscar(q: str = Query(..., min_length=1)):
    db = get_db()
    like = f"%{q}%"
    rows = db.execute("""
        SELECT a.id, a.nombre, s.cantidad as cantidad, a.unidad,
               s.ubicacion as ubicacion,
               p.id as proyecto_id, p.nombre as proyecto_nombre, p.color as proyecto_color,
               c.nombre as categoria_nombre, sal.nombre as sala_nombre
        FROM stock s
        JOIN articulos a ON a.id = s.articulo_id
        JOIN proyectos p ON p.id = s.proyecto_id
        LEFT JOIN categorias c ON c.id = a.categoria_id
        LEFT JOIN salas sal ON sal.id = s.sala_id
        WHERE s.estado = 'ok'
          AND (a.nombre LIKE ? OR a.marca LIKE ? OR a.referencia LIKE ?
           OR s.ubicacion LIKE ? OR c.nombre LIKE ? OR sal.nombre LIKE ?)
        ORDER BY a.nombre
        LIMIT 50
    """, (like, like, like, like, like, like)).fetchall()
    db.close()
    return [dict(r) for r in rows]


@app.get("/api/sugerencias", response_model=Sugerencias)
def get_sugerencias():
    """Valores existentes en el catálogo, para autocompletar nombre/marca/
    referencia/ubicación al dar de alta o editar artículos y evitar duplicados
    por pequeñas diferencias de escritura."""
    db = get_db()
    nombres = [r["nombre"] for r in db.execute("SELECT DISTINCT nombre FROM articulos ORDER BY nombre").fetchall()]
    marcas = [r["marca"] for r in db.execute("SELECT DISTINCT marca FROM articulos WHERE marca != '' ORDER BY marca").fetchall()]
    referencias = [r["referencia"] for r in db.execute("SELECT DISTINCT referencia FROM articulos WHERE referencia != '' ORDER BY referencia").fetchall()]
    ubicaciones = [r["ubicacion"] for r in db.execute("SELECT DISTINCT ubicacion FROM stock WHERE ubicacion IS NOT NULL ORDER BY ubicacion").fetchall()]
    db.close()
    return {"nombres": nombres, "marcas": marcas, "referencias": referencias, "ubicaciones": ubicaciones}


@app.get("/api/catalogo", response_model=List[CatalogoItem])
def buscar_catalogo(q: str = Query(..., min_length=1), excluir: Optional[int] = None):
    """Busca artículos del catálogo (no filas de stock) por nombre/marca/
    referencia — usado para elegir el artículo destino al fusionar duplicados."""
    db = get_db()
    like = f"%{q}%"
    rows = db.execute("""
        SELECT a.id, a.nombre, a.marca, a.referencia, c.nombre as categoria_nombre, a.unidad
        FROM articulos a
        LEFT JOIN categorias c ON c.id = a.categoria_id
        WHERE (a.nombre LIKE ? OR a.marca LIKE ? OR a.referencia LIKE ?)
          AND a.id != COALESCE(?, -1)
        ORDER BY a.nombre
        LIMIT 20
    """, (like, like, like, excluir)).fetchall()
    db.close()
    return [dict(r) for r in rows]


@app.post("/api/articulos/{origen_id}/fusionar")
def fusionar_articulos(origen_id: int, data: FusionArticulos):
    """Fusiona el artículo `origen_id` dentro de `articulo_destino_id`: mueve
    todas sus filas de stock (funcionales y de baja) al destino, sumando
    cantidades donde coincide proyecto+ubicación+estado y conservando las
    demás como ubicaciones separadas; reasigna el historial de movimientos;
    borra el artículo de origen. No se puede deshacer."""
    if origen_id == data.articulo_destino_id:
        raise HTTPException(400, "Selecciona un artículo distinto para fusionar")

    db = get_db()
    origen = db.execute("SELECT id FROM articulos WHERE id=?", (origen_id,)).fetchone()
    destino = db.execute("SELECT id FROM articulos WHERE id=?", (data.articulo_destino_id,)).fetchone()
    if not origen or not destino:
        db.close()
        raise HTTPException(404, "Artículo no encontrado")

    filas_origen = db.execute("SELECT * FROM stock WHERE articulo_id=?", (origen_id,)).fetchall()
    for fila in filas_origen:
        destino_stock_id = _find_or_create_stock(
            db, data.articulo_destino_id, fila["proyecto_id"], fila["ubicacion"], estado=fila["estado"]
        )
        db.execute(
            "UPDATE stock SET cantidad = cantidad + ?, updated_at=datetime('now') WHERE id=?",
            (fila["cantidad"], destino_stock_id)
        )
        if fila["sala_id"] is not None:
            destino_stock = db.execute("SELECT sala_id FROM stock WHERE id=?", (destino_stock_id,)).fetchone()
            if destino_stock["sala_id"] is None:
                db.execute("UPDATE stock SET sala_id=? WHERE id=?", (fila["sala_id"], destino_stock_id))
        db.execute("DELETE FROM stock WHERE id=?", (fila["id"],))

    db.execute("UPDATE movimientos SET articulo_id=? WHERE articulo_id=?", (data.articulo_destino_id, origen_id))
    db.execute("DELETE FROM articulos WHERE id=?", (origen_id,))
    db.commit()
    db.close()
    return {"detail": "ok", "articulo_id": data.articulo_destino_id}


# ─────────────────────────────────────────
# MOVIMIENTOS EN BULK
# ─────────────────────────────────────────

@app.post("/api/movimientos/bulk")
def create_movimientos_bulk(movimientos: List[MovimientoCreate]):
    db = get_db()
    try:
        for data in movimientos:
            stock = db.execute("SELECT * FROM stock WHERE id=?", (data.stock_id,)).fetchone()
            if not stock:
                raise HTTPException(404, f"Stock {data.stock_id} no encontrado")
            nueva = stock["cantidad"] + data.cantidad if data.tipo == "entrada" else stock["cantidad"] - data.cantidad
            if nueva < 0:
                articulo = db.execute("SELECT nombre FROM articulos WHERE id=?", (stock["articulo_id"],)).fetchone()
                raise HTTPException(400, f"Stock insuficiente para '{articulo['nombre']}' (actual: {stock['cantidad']})")
            db.execute(
                "INSERT INTO movimientos (articulo_id, tipo, cantidad, motivo, operador, proyecto_id) VALUES (?,?,?,?,?,?)",
                (stock["articulo_id"], data.tipo, data.cantidad, data.motivo, data.operador, stock["proyecto_id"])
            )
            if abs(nueva) < 1e-9:
                db.execute("DELETE FROM stock WHERE id=?", (data.stock_id,))
                otros = db.execute("SELECT COUNT(*) as n FROM stock WHERE articulo_id=?", (stock["articulo_id"],)).fetchone()
                if otros["n"] == 0:
                    db.execute("DELETE FROM articulos WHERE id=?", (stock["articulo_id"],))
            else:
                db.execute("UPDATE stock SET cantidad=?, updated_at=datetime('now') WHERE id=?", (nueva, data.stock_id))
        db.commit()
    finally:
        db.close()
    return {"actualizados": len(movimientos)}


# ─────────────────────────────────────────
# BACKUP / RESTORE
# ─────────────────────────────────────────

@app.get("/api/backup")
def backup_db():
    if not os.path.exists(DB_PATH):
        raise HTTPException(404, "Base de datos no encontrada")
    from datetime import datetime
    fecha = datetime.now().strftime("%Y%m%d_%H%M%S")
    return FileResponse(
        DB_PATH,
        media_type="application/octet-stream",
        headers={"Content-Disposition": f"attachment; filename=inventario_backup_{fecha}.db"}
    )

@app.post("/api/restore")
async def restore_db(file: UploadFile = File(...)):
    content = await file.read()
    if not content.startswith(b"SQLite format 3"):
        raise HTTPException(400, "El archivo no es una base de datos SQLite válida")
    with tempfile.NamedTemporaryFile(delete=False, suffix=".db") as tmp:
        tmp.write(content)
        tmp_path = tmp.name
    try:
        shutil.move(tmp_path, DB_PATH)
    except Exception as e:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)
        raise HTTPException(500, f"Error al restaurar: {e}")
    return {"message": "Base de datos restaurada. Recarga la aplicación."}


# ─────────────────────────────────────────
# EXPORTACIÓN EXCEL
# ─────────────────────────────────────────

@app.get("/api/proyectos/{proyecto_id}/exportar")
def exportar_excel(proyecto_id: int):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(500, "openpyxl no instalado")

    db = get_db()
    proyecto = db.execute("SELECT * FROM proyectos WHERE id=?", (proyecto_id,)).fetchone()
    if not proyecto:
        db.close()
        raise HTTPException(404, "Proyecto no encontrado")

    articulos = db.execute("""
        SELECT a.nombre, a.marca, a.referencia, c.nombre as categoria_nombre,
               SUM(s.cantidad) as cantidad, a.unidad,
               CASE WHEN COUNT(*) = 1 THEN MAX(s.ubicacion) ELSE 'Varias ubicaciones' END as ubicacion,
               a.stock_minimo as stock_minimo,
               MAX(s.updated_at) as updated_at
        FROM stock s
        JOIN articulos a ON a.id = s.articulo_id
        LEFT JOIN categorias c ON c.id = a.categoria_id
        WHERE s.proyecto_id = ? AND s.estado = 'ok'
        GROUP BY a.id
        ORDER BY a.nombre
    """, (proyecto_id,)).fetchall()

    movimientos = db.execute("""
        SELECT m.*, a.nombre as articulo_nombre
        FROM movimientos m
        JOIN articulos a ON a.id = m.articulo_id
        WHERE m.proyecto_id = ? OR m.proyecto_origen_id = ? OR m.proyecto_destino_id = ?
        ORDER BY m.fecha DESC
    """, (proyecto_id, proyecto_id, proyecto_id)).fetchall()
    db.close()

    wb = openpyxl.Workbook()

    # Hoja 1: Inventario
    ws1 = wb.active
    ws1.title = "Inventario"

    header_fill = PatternFill("solid", fgColor="1E3A5F")
    header_font = Font(bold=True, color="FFFFFF")
    headers = ["Nombre", "Marca", "Referencia", "Categoría", "Cantidad", "Unidad", "Ubicación", "Stock Mínimo", "Estado", "Última actualización"]

    for col, h in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, art in enumerate(articulos, 2):
        art = dict(art)
        bajo = art.get("stock_minimo") and art["cantidad"] < art["stock_minimo"]
        ws1.cell(row=row_idx, column=1, value=art["nombre"])
        ws1.cell(row=row_idx, column=2, value=art.get("marca") or "—")
        ws1.cell(row=row_idx, column=3, value=art.get("referencia") or "—")
        ws1.cell(row=row_idx, column=4, value=art.get("categoria_nombre") or "—")
        ws1.cell(row=row_idx, column=5, value=art["cantidad"])
        ws1.cell(row=row_idx, column=6, value=art.get("unidad") or "ud")
        ws1.cell(row=row_idx, column=7, value=art.get("ubicacion") or "—")
        ws1.cell(row=row_idx, column=8, value=art.get("stock_minimo") or "—")
        estado_cell = ws1.cell(row=row_idx, column=9, value="⚠ Bajo mínimo" if bajo else "OK")
        if bajo:
            estado_cell.font = Font(color="CC0000", bold=True)
        ws1.cell(row=row_idx, column=10, value=art.get("updated_at", ""))

    for col in range(1, len(headers) + 1):
        ws1.column_dimensions[get_column_letter(col)].width = 18

    # Hoja 2: Historial
    ws2 = wb.create_sheet("Historial de movimientos")
    headers2 = ["Fecha", "Artículo", "Tipo", "Cantidad", "Operador", "Motivo"]

    for col, h in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, mov in enumerate(movimientos, 2):
        mov = dict(mov)
        ws2.cell(row=row_idx, column=1, value=mov.get("fecha", ""))
        ws2.cell(row=row_idx, column=2, value=mov.get("articulo_nombre", ""))
        tipo_cell = ws2.cell(row=row_idx, column=3, value=mov["tipo"].capitalize())
        color_tipo = {"entrada": "006600", "salida": "CC0000", "transferencia": "1E3A5F", "baja": "B45309", "reparacion": "047857"}
        tipo_cell.font = Font(color=color_tipo.get(mov["tipo"], "000000"))
        ws2.cell(row=row_idx, column=4, value=mov["cantidad"])
        ws2.cell(row=row_idx, column=5, value=mov.get("operador", ""))
        ws2.cell(row=row_idx, column=6, value=mov.get("motivo") or "—")

    for col in range(1, len(headers2) + 1):
        ws2.column_dimensions[get_column_letter(col)].width = 20

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    nombre_archivo = f"inventario_{proyecto['nombre'].replace(' ', '_')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={nombre_archivo}"}
    )


# ─────────────────────────────────────────
# IMPORTACIÓN MASIVA DESDE EXCEL
# ─────────────────────────────────────────

IMPORT_COLUMNAS = ["proyecto", "nombre", "marca", "referencia", "categoria", "cantidad", "unidad", "ubicacion", "sala", "stock_minimo", "notas"]
IMPORT_CABECERAS = ["Proyecto", "Nombre", "Marca", "Referencia", "Categoría", "Cantidad", "Unidad", "Ubicación", "Sala", "Stock Mínimo", "Notas"]

def _celda_str(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None

def _celda_float(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        raise ValueError(f"'{v}' no es un número válido")

def _find_or_create_proyecto(db, nombre):
    row = db.execute("SELECT id FROM proyectos WHERE nombre=?", (nombre,)).fetchone()
    if row:
        return row["id"]
    cur = db.execute(
        "INSERT INTO proyectos (nombre, color, icono) VALUES (?,?,?)",
        (nombre, "#3B82F6", "Package")
    )
    return cur.lastrowid

def _procesar_importacion(db, content, proyecto_default_id):
    """Procesa las filas de un Excel de importación. `proyecto_default_id` es
    el proyecto a usar cuando la columna Proyecto de una fila viene vacía
    (el proyecto desde el que se sube, o 'Almacén general' en la importación
    global). Devuelve {procesadas, errores: [{fila, motivo}]}; las filas con
    error no frenan el resto."""
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(500, "openpyxl no instalado")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.worksheets[0]
    except Exception:
        raise HTTPException(400, "El archivo no es un Excel (.xlsx) válido")

    filas = list(ws.iter_rows(min_row=2, values_only=True))
    procesadas = 0
    errores = []

    for idx, fila in enumerate(filas, start=2):
        valores = list(fila or ())
        valores += [None] * (len(IMPORT_COLUMNAS) - len(valores))
        datos = dict(zip(IMPORT_COLUMNAS, valores))

        if all(_celda_str(v) is None for v in valores):
            continue  # fila vacía

        try:
            nombre = _celda_str(datos["nombre"])
            if not nombre:
                raise ValueError("El nombre es obligatorio")

            proyecto_nombre = _celda_str(datos["proyecto"])
            if proyecto_nombre:
                proyecto_id = _find_or_create_proyecto(db, proyecto_nombre)
            elif proyecto_default_id is not None:
                proyecto_id = proyecto_default_id
            else:
                raise ValueError("Falta la columna Proyecto y no hay un proyecto por defecto")

            cantidad = _celda_float(datos["cantidad"]) or 0
            stock_minimo = _celda_float(datos["stock_minimo"])
            marca = _celda_str(datos["marca"]) or ""
            referencia = _celda_str(datos["referencia"]) or ""
            unidad = _celda_str(datos["unidad"]) or "ud"
            ubicacion = _celda_str(datos["ubicacion"])
            notas = _celda_str(datos["notas"])

            categoria_id = None
            categoria_nombre = _celda_str(datos["categoria"])
            if categoria_nombre:
                db.execute("INSERT OR IGNORE INTO categorias (nombre) VALUES (?)", (categoria_nombre,))
                categoria_id = db.execute("SELECT id FROM categorias WHERE nombre=?", (categoria_nombre,)).fetchone()["id"]

            sala_id = None
            sala_nombre = _celda_str(datos["sala"])
            if sala_nombre:
                db.execute("INSERT OR IGNORE INTO salas (nombre) VALUES (?)", (sala_nombre,))
                sala_id = db.execute("SELECT id FROM salas WHERE nombre=?", (sala_nombre,)).fetchone()["id"]

            _crear_o_sumar_articulo(
                db, nombre=nombre, marca=marca, referencia=referencia,
                categoria_id=categoria_id, unidad=unidad, stock_minimo=stock_minimo,
                sala_id=sala_id, ubicacion=ubicacion, notas=notas,
                proyecto_id=proyecto_id, cantidad=cantidad,
                motivo="Importación Excel",
            )
            procesadas += 1
        except ValueError as e:
            errores.append({"fila": idx, "motivo": str(e)})
        except Exception as e:
            errores.append({"fila": idx, "motivo": f"Error inesperado: {e}"})

    return {"procesadas": procesadas, "errores": errores}

@app.get("/api/plantilla-importacion")
def plantilla_importacion():
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(500, "openpyxl no instalado")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Artículos"

    header_fill = PatternFill("solid", fgColor="1E3A5F")
    header_font = Font(bold=True, color="FFFFFF")
    for col, h in enumerate(IMPORT_CABECERAS, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col)].width = 18

    ejemplos = [
        ["Taller", "Tornillo M6x20", "Bosch", "TX-620", "Tornillería", 250, "ud", "Estantería A", "Taller", 50, ""],
        ["", "Cable eléctrico 2.5mm", "", "", "Cables y conectores", 30, "m", "Estantería B", "", "", "Rollo compartido"],
    ]
    for row_idx, fila in enumerate(ejemplos, 2):
        for col_idx, valor in enumerate(fila, 1):
            ws.cell(row=row_idx, column=col_idx, value=valor if valor != "" else None)

    ws2 = wb.create_sheet("Instrucciones")
    instrucciones = [
        "Cómo rellenar esta plantilla:",
        "",
        "- Proyecto: opcional, texto libre. Si se indica, la fila se importa a ese proyecto",
        "  (se crea si no existe todavía). Si se deja vacía, se usa el proyecto desde el que",
        "  subas el archivo, o 'Almacén general' si lo subes sin entrar en ningún proyecto.",
        "  Así puedes rellenar artículos de varios proyectos distintos en el mismo Excel.",
        "- Nombre: obligatorio.",
        "- Marca / Referencia: opcionales. Nombre + Marca + Referencia identifican el artículo:",
        "  si ya existe uno igual (en ese proyecto), se sumará la cantidad en vez de duplicarlo.",
        "- Categoría / Sala: opcionales, texto libre. Si no existen todavía se crean automáticamente.",
        "- Cantidad: opcional (si se deja vacía, se registra el artículo con 0 unidades).",
        "- Unidad: opcional, por defecto 'ud'.",
        "- Ubicación: opcional (Armario/Balda). Si dos filas del mismo artículo tienen ubicaciones",
        "  distintas, quedan como existencias separadas dentro del proyecto.",
        "- Stock Mínimo: opcional, numérico, para las alertas de stock bajo.",
        "- Notas: opcional.",
        "",
        "Las filas vacías se ignoran.",
    ]
    for row_idx, texto in enumerate(instrucciones, 1):
        ws2.cell(row=row_idx, column=1, value=texto)
    ws2.column_dimensions["A"].width = 95

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=plantilla_importacion_articulos.xlsx"}
    )

@app.post("/api/proyectos/{proyecto_id}/importar")
async def importar_excel(proyecto_id: int, file: UploadFile = File(...)):
    """Importa a un proyecto concreto. Las filas sin columna Proyecto van a
    este proyecto; las que sí la traigan pueden apuntar a otro (se crea si
    no existe)."""
    db = get_db()
    proyecto = db.execute("SELECT id FROM proyectos WHERE id=?", (proyecto_id,)).fetchone()
    if not proyecto:
        db.close()
        raise HTTPException(404, "Proyecto no encontrado")
    content = await file.read()
    try:
        resultado = _procesar_importacion(db, content, proyecto_default_id=proyecto_id)
        db.commit()
    finally:
        db.close()
    return resultado

@app.post("/api/importar")
async def importar_excel_global(file: UploadFile = File(...)):
    """Importación sin entrar en ningún proyecto: cada fila decide su destino
    con la columna Proyecto (se crea si no existe); si se deja vacía, va al
    Almacén general."""
    db = get_db()
    almacen = db.execute("SELECT id FROM proyectos WHERE es_almacen = 1").fetchone()
    content = await file.read()
    try:
        resultado = _procesar_importacion(
            db, content, proyecto_default_id=almacen["id"] if almacen else None
        )
        db.commit()
    finally:
        db.close()
    return resultado


# ─────────────────────────────────────────
# SERVIR FRONTEND (producción)
# ─────────────────────────────────────────

FRONTEND_DIST = os.path.join(os.path.dirname(__file__), "..", "frontend", "dist")

if os.path.exists(FRONTEND_DIST):
    app.mount("/", StaticFiles(directory=FRONTEND_DIST, html=True), name="static")
